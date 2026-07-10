#!/usr/bin/env python3
"""Convert 11 Beamer .tex slides to a PowerPoint presentation."""

import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SLIDES_DIR = "slides"
OUTPUT = "domain_partition_presentation.pptx"

# 16:9 aspect ratio
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

LEFT_MARGIN = Inches(0.6)
RIGHT_MARGIN = Inches(0.6)
TOP_MARGIN = Inches(0.8)


def strip_latex_formatting(text):
    """Remove formatting commands but keep math content."""
    text = re.sub(r"\\textbf\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\textit\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\emph\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\texttt\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\text\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\tiny", "", text)
    text = re.sub(r"\\footnotesize", "", text)
    text = re.sub(r"\\small", "", text)
    text = re.sub(r"\\Large", "", text)
    text = re.sub(r"\\large", "", text)
    text = re.sub(r"\\normalsize", "", text)
    text = re.sub(r"\\centering", "", text)
    text = re.sub(r"\\hfill", "  ", text)
    text = re.sub(r"\\quad", "  ", text)
    text = re.sub(r"\\qquad", "    ", text)
    text = re.sub(r"\\vspace\{[^}]*\}", "", text)
    text = re.sub(r"\\hspace\{[^}]*\}", "", text)
    text = re.sub(r"\\left", "", text)
    text = re.sub(r"\\right", "", text)
    text = re.sub(r"\\bigl", "", text)
    text = re.sub(r"\\bigr", "", text)
    text = re.sub(r"\\big", "", text)
    text = re.sub(r"\\!", "", text)
    text = re.sub(r"\\,", "", text)
    text = re.sub(r"\\;", "", text)
    text = re.sub(r"\\:", "", text)
    text = re.sub(r"\\&", "&", text)
    text = re.sub(r"\\#", "#", text)
    text = re.sub(r"\\%", "%", text)
    text = re.sub(r"\\$", "$", text)
    text = re.sub(r"\\_", "_", text)
    text = re.sub(r"  +", " ", text)
    return text.strip()


def parse_tex_slide(filepath):
    """Parse a Beamer .tex file and extract title, bullets, equations, citation."""
    with open(filepath, "r", encoding="utf-8") as f:
        content = f.read()

    frame_pattern = r"\\begin\{frame\}(?:\{([^}]*)\})?(.*?)\\end\{frame\}"
    frames = re.findall(frame_pattern, content, re.DOTALL)

    content_frames = []
    for title, body in frames:
        if title and title.strip():
            content_frames.append((title.strip(), body))

    if not content_frames:
        return None

    title, body = content_frames[0]

    bullets = []
    itemize_pattern = r"\\begin\{itemize\}(.*?)\\end\{itemize\}"
    itemizes = re.findall(itemize_pattern, body, re.DOTALL)
    for itemize_body in itemizes:
        items = re.findall(r"\\item\s*(.*)", itemize_body)
        for item in items:
            item = item.strip()
            if item:
                bullets.append(strip_latex_formatting(item))

    equations = []
    eq_pattern = r"\\\[(.*?)\\\]"
    eqs = re.findall(eq_pattern, body, re.DOTALL)
    for eq in eqs:
        eq = eq.strip()
        if eq:
            equations.append(strip_latex_formatting(eq))

    inline_pattern = r"\$([^$]+)\$"
    inlines = re.findall(inline_pattern, body)
    for inl in inlines:
        inl = inl.strip()
        if inl and len(inl) > 5:
            equations.append(strip_latex_formatting(inl))

    citation = ""
    cit_patterns = [
        r"Citation:\s*(.*)",
        r"\\textit\{Lit:\s*([^}]*)\}",
        r"\\textit\{Ref:\s*([^}]*)\}",
        r"Lit:\s*(.*)",
        r"Ref:\s*(.*)",
    ]
    for pattern in cit_patterns:
        match = re.search(pattern, body)
        if match:
            citation = strip_latex_formatting(match.group(1)).strip()
            if citation:
                break

    if not citation:
        lines = body.split("\n")
        for line in reversed(lines):
            line = line.strip()
            if line.startswith("Citation:") or line.startswith("Lit:") or line.startswith("Ref:"):
                citation = strip_latex_formatting(line.split(":", 1)[1]).strip()
                if citation:
                    break

    return {
        "title": strip_latex_formatting(title),
        "bullets": bullets,
        "equations": equations,
        "citation": citation,
    }


def add_slide(prs, data):
    """Add a slide to the presentation."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    title_box = slide.shapes.add_textbox(
        LEFT_MARGIN, TOP_MARGIN - Inches(0.3),
        SLIDE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN, Inches(0.8)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = data["title"]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    p.alignment = PP_ALIGN.LEFT

    content_top = TOP_MARGIN + Inches(0.6)
    content_height = SLIDE_HEIGHT - content_top - Inches(0.8)

    if data["equations"]:
        eq_height = min(Inches(1.2), Inches(0.4 * len(data["equations"])))
        bullet_height = content_height - eq_height - Inches(0.3)
    else:
        eq_height = Inches(0)
        bullet_height = content_height

    if data["bullets"]:
        bullet_box = slide.shapes.add_textbox(
            LEFT_MARGIN, content_top,
            SLIDE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN, bullet_height
        )
        tf = bullet_box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(data["bullets"]):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(20)
            p.font.color.rgb = RGBColor(0x30, 0x30, 0x30)
            p.level = 0
            p.space_after = Pt(12)

    if data["equations"]:
        eq_top = content_top + bullet_height + Inches(0.2)
        eq_box = slide.shapes.add_textbox(
            LEFT_MARGIN, eq_top,
            SLIDE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN, eq_height
        )
        tf = eq_box.text_frame
        tf.word_wrap = True
        for i, eq in enumerate(data["equations"]):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = eq
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            p.font.italic = True
            p.alignment = PP_ALIGN.CENTER
            p.space_after = Pt(8)

    if data["citation"]:
        cit_box = slide.shapes.add_textbox(
            LEFT_MARGIN, SLIDE_HEIGHT - Inches(0.7),
            SLIDE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN, Inches(0.4)
        )
        tf = cit_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"Citation: {data['citation']}"
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        p.font.italic = True
        p.alignment = PP_ALIGN.LEFT

    return slide


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slide_files = sorted([
        f for f in os.listdir(SLIDES_DIR)
        if f.startswith("slide") and f.endswith(".tex")
    ])

    print(f"Found {len(slide_files)} slide files")

    for filename in slide_files:
        filepath = os.path.join(SLIDES_DIR, filename)
        data = parse_tex_slide(filepath)
        if data:
            add_slide(prs, data)
            print(f"  Added: {filename} -> {data['title']}")
        else:
            print(f"  Skipped (no content frame): {filename}")

    prs.save(OUTPUT)
    print(f"\nSaved {len(prs.slides)} slides to {OUTPUT}")
    print(f"Slide dimensions: {prs.slide_width} x {prs.slide_height}")


if __name__ == "__main__":
    main()
